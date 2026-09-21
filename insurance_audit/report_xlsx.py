"""엑셀 보고서 산출.

보고와 공유는 결국 엑셀로 이뤄지므로 시트를 목적별로 나눠 둔다.
순위 시트에서 혐의자를 고르고, 지표 시트에서 근거를 확인하고,
소명 시트에서 건별 목록을 뽑는 흐름이다.
"""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd

from . import config
from .metrics.base import MetricResult

HEADER_FILL = "FFD9E1F2"
BAND_FILL = {"높음": "FFF8CBAD", "중간": "FFFFE699", "낮음": None}

# openpyxl 이 거부하는 제어문자. 값 하나만 섞여 있어도 저장 전체가 실패하므로
# 분석을 다 끝내고 결과를 잃는다. 시트로 넘기기 직전에 걷어낸다.
_ILLEGAL = r"[\x00-\x08\x0b-\x0c\x0e-\x1f\x7f]"


def _sheet(writer, frame: pd.DataFrame, name: str, index: bool = False) -> None:
    """제어문자를 걷어내고 시트로 쓴다. 시트 이름 길이 제한도 여기서 맞춘다."""
    safe = frame.copy()
    for column in safe.columns:
        if safe[column].dtype == object or safe[column].dtype == "string":
            safe[column] = safe[column].astype("string").str.replace(
                _ILLEGAL, "", regex=True
            )
    safe.columns = [re.sub(_ILLEGAL, "", str(c)) for c in safe.columns]
    if index:
        # 업체×연도 행렬처럼 이름이 인덱스로 들어가는 시트가 있다.
        safe.index = [re.sub(_ILLEGAL, "", str(v)) for v in safe.index]
    safe.to_excel(writer, sheet_name=name[:31], index=index)


def write(
    path: str | Path,
    rankings: dict[str, pd.DataFrame],
    results: list[MetricResult],
    data: dict[str, pd.DataFrame],
    drill: pd.DataFrame,
    column_check: pd.DataFrame,
    views: dict[str, object] | None = None,
) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    views = views or {}

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        _sheet(writer, _summary(data, results, rankings), "00_요약")

        for order, entity in enumerate(("담당자", "결재자", "법인", "소속", "조합"), start=1):
            frame = rankings.get(entity)
            if frame is None or frame.empty:
                continue
            _sheet(writer, frame, f"{order:02d}_{entity}순위")

        _write_views(writer, views)

        for order, result in enumerate(results, start=20):
            if result.detail is None or result.detail.empty:
                continue
            _sheet(writer, result.detail, f"{order}_{result.name}")

        if not drill.empty:
            _sheet(writer, drill, "90_소명대상건")
        if not column_check.empty:
            _sheet(writer, column_check, "99_컬럼점검")

        _style(writer)

    return path


def _write_views(writer: pd.ExcelWriter, views: dict[str, object]) -> None:
    """현황 표. 혐의 순위와 별개로 전체 그림을 보는 시트다."""
    sheets = (
        ("10_연도별업체순위", "연도별업체"),
        ("11_업체×연도위임율", "업체연도행렬"),
        ("12_소속2별업체", "소속2별업체"),
        ("13_소속3별업체", "소속3별업체"),
        ("14_담당자별업체", "담당자별업체"),
    )
    for sheet, key in sheets:
        frame = views.get(key)
        if not isinstance(frame, pd.DataFrame) or frame.empty:
            continue
        _sheet(writer, frame, sheet, index=key == "업체연도행렬")


def _summary(
    data: dict[str, pd.DataFrame],
    results: list[MetricResult],
    rankings: dict[str, pd.DataFrame],
) -> pd.DataFrame:
    cases = data["cases"]
    payments = data["payments"]

    rows = [
        {"구분": "데이터", "항목": "지급 행 수", "값": f"{len(payments):,}"},
        {"구분": "데이터", "항목": "조사 배당 건수", "값": f"{len(cases):,}"},
        {"구분": "데이터", "항목": "담당자 수", "값": f"{cases['담당자_id'].nunique():,}"},
        {"구분": "데이터", "항목": "손사법인 수", "값": f"{cases['법인_id'].nunique():,}"},
    ]
    if "결재자_id" in cases.columns:
        rows.append(
            {"구분": "데이터", "항목": "결재자 수", "값": f"{cases['결재자_id'].nunique():,}"}
        )
    if "사고일" in cases.columns and cases["사고일"].notna().any():
        rows.append(
            {
                "구분": "데이터",
                "항목": "사고일 범위",
                "값": f"{cases['사고일'].min():%Y-%m-%d} ~ {cases['사고일'].max():%Y-%m-%d}",
            }
        )

    for result in results:
        status = "산출" if result.scores else "미산출"
        rows.append(
            {
                "구분": "지표",
                "항목": f"{result.name} ({result.title})",
                "값": f"{status} · {result.note}" if result.note else status,
            }
        )

    for entity, frame in rankings.items():
        if frame.empty:
            continue
        high = int((frame["등급"] == config.RISK_RULES[0][2]).sum())
        rows.append(
            {
                "구분": "결과",
                "항목": f"{entity} 대상 수",
                "값": f"{len(frame):,}명/곳 (높음 {high}건)",
            }
        )

    rows.append(
        {
            "구분": "참고",
            "항목": "해석 주의",
            "값": "혐의도는 통계적 이상 신호이며 그 자체가 비위 증거가 아님. 소명 절차로 확인 필요.",
        }
    )
    return pd.DataFrame(rows)


def _style(writer: pd.ExcelWriter) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True)

    for sheet in writer.book.worksheets:
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        band_column = _find_column(sheet, "등급")
        if band_column:
            for row in range(2, sheet.max_row + 1):
                value = sheet.cell(row=row, column=band_column).value
                color = BAND_FILL.get(str(value))
                if color:
                    for column in range(1, sheet.max_column + 1):
                        sheet.cell(row=row, column=column).fill = PatternFill(
                            "solid", fgColor=color
                        )

        for column_cells in sheet.columns:
            longest = max(
                (len(str(cell.value)) for cell in column_cells if cell.value is not None),
                default=8,
            )
            letter = column_cells[0].column_letter
            sheet.column_dimensions[letter].width = min(max(longest + 2, 10), 55)


def _find_column(sheet, header: str) -> int | None:
    for cell in sheet[1]:
        if cell.value == header:
            return cell.column
    return None
